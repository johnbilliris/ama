"""Quick inspector for the AMA pptx template & example."""
import sys
from pptx import Presentation

target = sys.argv[1] if len(sys.argv) > 1 else "docs/ama-presentation/Azure Master Architect_Template.potm"

p = Presentation(target)
print("File:", target)
print("Slides:", len(p.slides))
print("Slide size (EMU):", p.slide_width, "x", p.slide_height)
print("Slide masters:", len(p.slide_masters))
for sm in p.slide_masters:
    print(f"  Master layouts ({len(sm.slide_layouts)}):")
    for i, layout in enumerate(sm.slide_layouts):
        ph_kinds = []
        for ph in layout.placeholders:
            ph_kinds.append(f"idx={ph.placeholder_format.idx}/type={ph.placeholder_format.type}/name={ph.name}")
        print(f"    [{i}] {layout.name} -> {ph_kinds}")

print()
print("=== Slides in this file ===")
for i, slide in enumerate(p.slides):
    print(f"--- Slide {i+1} (layout: {slide.slide_layout.name}) ---")
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            print(f"  NOTES: {notes[:600]}")
    for ph in slide.placeholders:
        try:
            t = ph.text_frame.text.strip() if ph.has_text_frame else ""
        except Exception:
            t = ""
        if t:
            print(f"  PH idx={ph.placeholder_format.idx} ({ph.name}): {t[:300]}")
